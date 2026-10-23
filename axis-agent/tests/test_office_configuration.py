"""Office configuration sources must survive extraction and planner pagination."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook

from axis.attachments.models import DocumentRequest
from axis.attachments.parsing import parse
from axis.attachments.runtime import DocumentSession
from test_attachments import documents, ingest, link
from test_attachment_workflows import instruction, workbook, followup


def test_word_nested_configuration_table_and_content_control(tmp_path):
    doc = Document()
    doc.add_heading('Production', 1)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = 'Configure service'
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = 'Retry limit'
    nested.cell(0, 1).text = '7'
    cell.add_paragraph('Save after setting retries')
    paragraph = doc.add_paragraph('Enable audit logging')
    control = OxmlElement('w:sdt')
    content = OxmlElement('w:sdtContent')
    paragraph._p.addprevious(control)
    control.append(content)
    content.append(paragraph._p)
    doc.add_paragraph('Verify audit logging')
    path = tmp_path / 'configuration.docx'
    doc.save(path)
    sections, _ = parse(path)
    text = '\n'.join(s['text'] for s in sections)
    assert 'Retry limit' in text and '7' in text
    assert text.index('Retry limit') < text.index('Save after setting retries')
    assert text.index('Enable audit logging') < text.index('Verify audit logging')


def test_excel_reads_cells_beyond_incorrect_dimension(tmp_path):
    book = Workbook()
    book.active.append(['Setting', 'Value'])
    book.active.append(['Timeout', 30])
    book.active.append(['Retries', 7])
    output = BytesIO()
    book.save(output)
    path = tmp_path / 'config.xlsx'
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(path, 'w', ZIP_DEFLATED) as dest:
        for name in source.namelist():
            data = source.read(name)
            if name == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension ref="[^"]+"/>', b'<dimension ref="A1:A1"/>', data)
            dest.writestr(name, data)
    sections, _ = parse(path)
    assert [s['row_number'] for s in sections] == [1, 2, 3]
    assert sections[-1]['cells'][1]['value'] == 7


def test_workflow_source_context_advances_past_planned_sections(documents):
    files, cid, task = documents
    item = ingest(files, cid, 'procedure.txt', '\n'.join(f'Step {i}: set value {i}' for i in range(40)).encode())
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    for _ in range(3):
        session.execute(DocumentRequest(operation='enumerate', attachment_id=item['id']))
    first = session.context()['procedure_sources']
    session.execute(DocumentRequest(operation='skip', section_ids=[s['section_id'] for s in first[:16]], reason='Reference only'))
    session.execute(DocumentRequest(operation='skip', section_ids=[s['section_id'] for s in first[16:]], reason='Reference only'))
    packet = DocumentSession(files, task).context()
    assert packet['procedure_sources'][0]['text'] == 'Step 24: set value 24'
    assert packet['procedure_sources'][-1]['text'] == 'Step 39: set value 39'
    assert packet['procedure_sources_has_more'] is False


def test_workflow_source_preserves_long_configuration_values(documents):
    files, cid, task = documents
    text = 'Configure: ' + 'x' * 850 + ' timeout=37'
    item = ingest(files, cid, 'procedure.txt', text.encode())
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    session.execute(DocumentRequest(operation='read', attachment_id=item['id']))
    assert session.context()['procedure_sources'][0]['text'] == text


def test_scoped_excel_configuration_is_present_in_procedure_context(documents):
    files, cid, task = documents
    item = workbook(files, cid, task)
    prompt = 'Apply the configuration in Test Results!A2:B3.'
    instruction(files, task, prompt)
    session = DocumentSession(files, task)
    session.execute(DocumentRequest(operation='read', attachment_id=item['id']))
    sources = session.context()['procedure_sources']
    assert len(sources) == 2
    assert sources[0]['cells'][0]['address'] == 'A2'
    assert 'C2' not in str(sources)


def test_narrowed_workflow_context_does_not_reexpose_other_columns(documents):
    files, cid, task = documents
    item = workbook(files, cid, task)
    prompt = 'Apply all configurations in the file.'
    instruction(files, task, prompt)
    session = DocumentSession(files, task)
    session.execute(DocumentRequest(operation='read', attachment_id=item['id']))
    followup(files, cid, task, 'Now apply only Test Results!A2:B3.')
    session.refresh_scope()
    sources = session.context()['procedure_sources']
    assert len(sources) == 2
    assert 'C2' not in str(sources)


def test_resumed_step_restores_exact_source_without_another_read(documents):
    from axis.attachments.models import WorkflowItem
    files, cid, task = documents
    value = 'endpoint=' + 'x' * 800 + '.example.test'
    item = ingest(files, cid, 'config.txt', value.encode())
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    source = session.execute(DocumentRequest(operation='read', attachment_id=item['id']))['items'][0]
    session.execute(DocumentRequest(operation='plan', steps=[WorkflowItem(
        section_id=source['id'], instruction='Configure endpoint', expected_result='Saved endpoint matches source')]))
    step = session.workflow()['next_steps'][0]
    session.begin(step['id'], 'goal_1')
    recovered = DocumentSession(files, task)
    assert recovered.last_result is None
    assert recovered.current_step is None
    assert recovered.context()['current_source']['items'][0]['text'] == value
    assert recovered.context()['workflow']['next_steps'][0]['goal_id'] == 'goal_1'


def test_workflow_context_is_bounded_and_does_not_expose_unread_columns(documents):
    import json
    files, cid, task = documents
    item = workbook(files, cid, task)
    instruction(files, task, 'Apply all configurations in the file.')
    session = DocumentSession(files, task)
    session.execute(DocumentRequest(operation='read', attachment_id=item['id'], sheet='Test Results', cell_range='A2:B3'))
    # A partial range read does not authorize presenting the entire row as read.
    assert session.context()['procedure_sources'] == []
    # Once the whole source has been read it can be retained, without another model read.
    session.execute(DocumentRequest(operation='read', attachment_id=item['id']))
    sources = session.context()['procedure_sources']
    assert sources and len(json.dumps(sources)) < 20000


def test_long_workflow_packet_advertises_remaining_sources(documents):
    files, cid, task = documents
    data = '\n'.join(f'Step {i}: ' + 'x' * 900 for i in range(40))
    item = ingest(files, cid, 'long.txt', data.encode())
    link(files, cid, task, item, 'instructions')
    session = DocumentSession(files, task)
    while not session.coverage()[0]['complete']:
        session.execute(DocumentRequest(operation='enumerate', attachment_id=item['id']))
    packet = session.context()
    assert len(packet['procedure_sources']) < 24
    assert packet['procedure_sources_has_more'] is True
    assert all(len(source['text']) > 900 for source in packet['procedure_sources'])
